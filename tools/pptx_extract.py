#!/usr/bin/env python3
"""Take a .pptx apart so an agent can rebuild it as a deck.

Usage:
    pip install python-pptx
    python3 tools/pptx_extract.py path/to/deck.pptx            # writes source/ and deck/assets/

Output:
    source/slides.md     one section per slide: title, text (with bullet levels), speaker notes,
                         tables, chart data and the images saved from that slide. Read this first.
    source/slides.json   the same, structured.
    deck/assets/sNN-K.*  every picture, named by slide number.

Nothing is guessed or summarised here; the agent does the rewriting.
"""
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Emu
except ImportError:
    sys.exit("python-pptx is missing. Install it with:  pip install python-pptx")

ROOT = Path(__file__).resolve().parent.parent


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def chart_data(chart):
    out = {"type": str(chart.chart_type).split(".")[-1].split(" ")[0], "categories": [], "series": []}
    try:
        plot = chart.plots[0]
        out["categories"] = [str(c) for c in plot.categories]
        for s in plot.series:
            out["series"].append({"name": s.name, "values": [None if v is None else float(v) for v in s.values]})
    except Exception as e:  # charts with exotic layouts still give us their title
        out["error"] = str(e)
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            out["title"] = chart.chart_title.text_frame.text
    except Exception:
        pass
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1]).expanduser()
    prs = Presentation(str(src))
    out_dir, assets = ROOT / "source", ROOT / "deck" / "assets"
    out_dir.mkdir(exist_ok=True)
    assets.mkdir(parents=True, exist_ok=True)
    slides = []
    for n, slide in enumerate(prs.slides, 1):
        s = {"slide": n, "layout": slide.slide_layout.name, "title": "", "text": [], "notes": "",
             "tables": [], "charts": [], "images": []}
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            s["title"] = slide.shapes.title.text_frame.text.strip()
        k = 0
        shapes = sorted(walk(slide.shapes), key=lambda x: ((x.top or 0), (x.left or 0)))
        for sh in shapes:
            if sh.has_text_frame and sh != slide.shapes.title:
                for p in sh.text_frame.paragraphs:
                    line = "".join(r.text for r in p.runs).strip()
                    if line:
                        s["text"].append({"level": p.level, "text": line})
            if getattr(sh, "has_table", False) and sh.has_table:
                s["tables"].append([[c.text.strip() for c in row.cells] for row in sh.table.rows])
            if getattr(sh, "has_chart", False) and sh.has_chart:
                s["charts"].append(chart_data(sh.chart))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(sh, "image", None) is not None:
                try:
                    img = sh.image
                except Exception:
                    continue
                k += 1
                name = f"s{n:02d}-{k}.{img.ext}"
                (assets / name).write_bytes(img.blob)
                s["images"].append({"file": f"assets/{name}", "width_px": int(Emu(sh.width).inches * 96) if sh.width else None})
        if slide.has_notes_slide:
            s["notes"] = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(s)

    (out_dir / "slides.json").write_text(json.dumps(slides, ensure_ascii=False, indent=1), encoding="utf-8")
    md = [f"# {src.name} — {len(slides)} slides\n", "Images are saved under deck/assets/ (paths below are relative to deck/).\n"]
    for s in slides:
        md.append(f"\n## Slide {s['slide']}: {s['title'] or '(no title)'}\n")
        md.append(f"_layout: {s['layout']}_\n")
        for t in s["text"]:
            md.append("  " * t["level"] + f"- {t['text']}")
        for i, tb in enumerate(s["tables"], 1):
            md.append(f"\n**Table {i}**\n")
            if tb:
                md.append("| " + " | ".join(tb[0]) + " |")
                md.append("|" + " --- |" * len(tb[0]))
                for row in tb[1:]:
                    md.append("| " + " | ".join(row) + " |")
        for i, ch in enumerate(s["charts"], 1):
            md.append(f"\n**Chart {i}** ({ch.get('type')}){': ' + ch['title'] if ch.get('title') else ''}")
            md.append(f"- categories: {ch.get('categories')}")
            for se in ch.get("series", []):
                md.append(f"- {se['name']}: {se['values']}")
        for im in s["images"]:
            md.append(f"\n![slide {s['slide']} image]({im['file']})")
        if s["notes"]:
            md.append("\n> speaker notes:\n> " + s["notes"].replace("\n", "\n> "))
    (out_dir / "slides.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    imgs = sum(len(s["images"]) for s in slides)
    print(f"{len(slides)} slides -> source/slides.md, source/slides.json; {imgs} images -> deck/assets/")


if __name__ == "__main__":
    main()
